#!/usr/bin/env python
"""Structural + numerical self-check for the standards archive.

    python scripts/verify.py          # structure + numbers (fast)
    python scripts/verify.py --hash   # additionally re-hash every file against MANIFEST.csv
"""
import csv, hashlib, os, sys, tarfile, zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FAIL = []
OK = []

def check(cond, msg):
    (OK if cond else FAIL).append(msg)
    print(("  PASS  " if cond else "  FAIL  ") + msg)

def size_of(rel):
    p = os.path.join(ROOT, rel.replace("/", os.sep))
    return os.path.getsize(p) if os.path.exists(p) else -1

print("== 1. required files")
REQUIRED = [
 "README.md", "HVAB/README_HVAB.md", "TN2474_Castles_Gray/README_TN2474.md",
 "HVAB/cases_hvab.csv", "HVAB/run_matrix.csv", "HVAB/10_papers/INDEX.md",
 "HVAB/01_general_information/HVAB_General_Information_Readme_v2.pdf",
 "HVAB/02_performance/Runs92_95_v2.1.xlsx", "HVAB/06_piv/VectorFielddata.zip",
 "HVAB/07_pressure_airloads/Run72_Airloads_v2.xlsx",
 "HVAB/08_workshop_HPW/fileshare_geometry/HVAB_Geometry/HVAB_UndeformedSurfAndVol.tar",
 "HVAB/09_nasa_docs/NASA-TM-2020-5002153_HVAB_Blade_Set_Definition.pdf",
 "TN2474_Castles_Gray/naca-tn-2474.pdf",
]
for rel in REQUIRED:
    check(size_of(rel) > 0, "%s (%d bytes)" % (rel, size_of(rel)))

print("\n== 2. performance truth table")
cases = list(csv.DictReader(open(os.path.join(ROOT, "HVAB", "cases_hvab.csv"), encoding="utf-8")))
check(len(cases) >= 130, "cases_hvab.csv rows = %d (expect ~137)" % len(cases))
bad_fm, bad_sig, bad_nan = [], [], []
for r in cases:
    ct, cp, fm = float(r["CT"]), float(r["CP"]), float(r["FMERIT"])
    if abs(fm - ct ** 1.5 / (2 ** 0.5 * cp)) > 2e-3:
        bad_fm.append(r["case"])
    sig_ct, sig_cp = ct / float(r["CTOS"]), cp / float(r["CPOS"])
    if abs(sig_ct - sig_cp) > 1e-4 or abs(sig_ct - 0.1033) > 1e-3:
        bad_sig.append(r["case"])
    if not (0 < ct < 0.02 and 0 < cp < 0.003 and 0 < fm < 0.85):
        bad_nan.append(r["case"])
check(not bad_fm, "FM = CT^1.5/(sqrt(2)*CP) for all rows (%s)" % (bad_fm[:4] or "ok"))
check(not bad_sig, "CT/CTOS = CP/CPOS = sigma = 0.1033 (%s)" % (bad_sig[:4] or "ok"))
check(not bad_nan, "CT/CP/FM inside physical ranges (%s)" % (bad_nan[:4] or "ok"))

declared = set()
for r in csv.DictReader(open(os.path.join(ROOT, "HVAB", "run_matrix.csv"), encoding="utf-8")):
    for x in r["runs"].split(","):
        declared.add(int(x))
have = {int(r["run"]) for r in cases}
check(declared == have, "run coverage matches run_matrix.csv (%d runs)" % len(declared))

print("\n== 3. archives")
z = zipfile.ZipFile(os.path.join(ROOT, "HVAB", "06_piv", "VectorFielddata.zip"))
check(z.testzip() is None and len(z.namelist()) > 100, "PIV VectorFielddata.zip valid, %d entries" % len(z.namelist()))
t = tarfile.open(os.path.join(ROOT, "HVAB", "08_workshop_HPW", "fileshare_geometry", "HVAB_Geometry",
                              "HVAB_UndeformedSurfAndVol.tar"), mode="r|")
names = [m.name for m in t]
t.close()
check(len(names) > 0, "HVAB_UndeformedSurfAndVol.tar readable, %d members: %s" % (len(names), names[:4]))

print("\n== 4. spreadsheets")
import openpyxl
for rel in ["HVAB/02_performance/Run30_v2.xlsx", "HVAB/02_performance/Runs92_95_v2.1.xlsx",
            "HVAB/07_pressure_airloads/Run72_Airloads_v2.xlsx", "HVAB/04_shadowgraphy/R50-52-54-SG-Data_24Sep2023.xlsx",
            "HVAB/05_photogrammetry/Run30_1250_BD.xlsx"]:
    p = os.path.join(ROOT, rel.replace("/", os.sep))
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        check(len(wb.sheetnames) > 0, "%s sheets=%s" % (rel, wb.sheetnames))
        wb.close()
    except Exception as e:
        check(False, "%s -> %s" % (rel, e))

print("\n== 5. pdfs")
from pypdf import PdfReader
for rel, minp in [("TN2474_Castles_Gray/naca-tn-2474.pdf", 70),
                  ("HVAB/09_nasa_docs/NASA-TM-2020-5002153_HVAB_Blade_Set_Definition.pdf", 40),
                  ("HVAB/09_nasa_docs/NASA_TM_HVAB_C81_Deck_Generation_2025.pdf", 15),
                  ("HVAB/02_performance/Performance_ReadMe_v2.1.pdf", 5),
                  ("HVAB/10_papers/Sheng_HybridRANSLES_InstalledRotor_ICCFD12.pdf", 15)]:
    try:
        n = len(PdfReader(os.path.join(ROOT, rel.replace("/", os.sep))).pages)
        check(n >= minp, "%s pages=%d" % (rel, n))
    except Exception as e:
        check(False, "%s -> %s" % (rel, e))

if "--hash" in sys.argv:
    print("\n== 6. sha256 vs MANIFEST.csv")
    man = os.path.join(ROOT, "MANIFEST.csv")
    rows = list(csv.DictReader(open(man, encoding="utf-8")))
    bad = 0
    for r in rows:
        p = os.path.join(ROOT, r["relpath"].replace("/", os.sep))
        if not os.path.exists(p):
            bad += 1; print("  MISSING", r["relpath"]); continue
        h = hashlib.sha256()
        with open(p, "rb") as f:
            for c in iter(lambda: f.read(1 << 20), b""):
                h.update(c)
        if h.hexdigest() != r["sha256"]:
            bad += 1; print("  MISMATCH", r["relpath"])
    check(bad == 0, "%d files hashed, %d mismatches" % (len(rows), bad))

print("\n%s: %d passed, %d failed" % ("RESULT", len(OK), len(FAIL)))
sys.exit(1 if FAIL else 0)
