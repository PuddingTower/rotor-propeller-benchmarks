import csv, glob, os, re
import openpyxl

ROOT = r"C:\翼型设计\标准算例"
PERF = os.path.join(ROOT, "HVAB", "02_performance")
OUT = os.path.join(ROOT, "HVAB", "cases_hvab.csv")

def sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(v is None for v in r):
            continue
        out.append({hdr[j]: r[j] for j in range(min(len(hdr), len(r))) if hdr[j]})
    return out

rows = []
for p in sorted(glob.glob(os.path.join(PERF, "Run*.xlsx"))):
    run_file = os.path.basename(p)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    avg = {r["Name"]: r for r in sheet_rows(wb["averages"]) if r.get("Name")}
    err = {r["Name"]: r for r in sheet_rows(wb["errors"]) if r.get("Name")} if "errors" in wb.sheetnames else {}
    for name, a in avg.items():
        e = err.get(name, {})
        m = re.match(r"R(\d+)M(\d+)TH(\d+)", str(name))
        rec = {
            "case": name,
            "run": int(m.group(1)) if m else "",
            "mtip_nom": int(m.group(2)) / 1000 if m else "",
            "collective_deg": int(m.group(3)) if m else "",
            "source_file": run_file,
        }
        for k, v in a.items():
            if k == "Name":
                continue
            rec[k] = v
        for k, v in e.items():
            if k in ("Name", "COLL"):
                continue
            rec[k + "_err"] = v
        rows.append(rec)

cols = ["case", "run", "mtip_nom", "collective_deg", "source_file",
        "RPM", "MTIP", "COLL_SP", "PS", "TSF", "RH", "RHO", "CSND",
        "BPITCH", "FLAP", "LAG", "THRUST", "TORQ", "CT", "CTOS", "CP", "CPOS", "CPOS ID", "FMERIT",
        "CT_err", "CP_err", "FMERIT_err", "THRUST_err", "TORQ_err", "MTIP_err", "RPM_err"]
cols = cols + [c for c in rows[0].keys() if c not in cols]
with open(OUT, "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    for r in sorted(rows, key=lambda r: (str(r["mtip_nom"]), r["run"], r["collective_deg"])):
        w.writerow(r)

print("cases: %d -> %s" % (len(rows), OUT))
print()
print("%-16s %6s %6s %6s %9s %9s %9s %8s" % ("case", "MTIP", "COLL", "RPM", "CT", "CP", "FM", "THRUST"))
for r in sorted(rows, key=lambda r: (str(r["mtip_nom"]), r["run"], r["collective_deg"])):
    print("%-16s %6.3f %6.1f %7.1f %9.5f %9.6f %9.4f %8.1f" % (
        r["case"], r.get("MTIP") or 0, r.get("COLL_SP") or 0, r.get("RPM") or 0,
        r.get("CT") or 0, r.get("CP") or 0, r.get("FMERIT") or 0, r.get("THRUST") or 0))
# self-check: FM must equal CT^1.5/(sqrt(2)*CP)
bad = [r["case"] for r in rows if r.get("CT") and r.get("CP") and r.get("FMERIT")
       and abs(r["FMERIT"] - r["CT"] ** 1.5 / (2 ** 0.5 * r["CP"])) > 2e-3]
print("\nFM self-check failures:", bad or "none")
